from pptx import Presentation

prs = Presentation('presentation/Context_Engineering_y_Ambiguedad_final_matematica.pptx')
for i, slide in enumerate(prs.slides):
    text = "No text"
    for shape in slide.shapes:
        if shape.has_text_frame:
            text = shape.text
            break
    print(f"Slide {i+1}: {text[:50]}")
