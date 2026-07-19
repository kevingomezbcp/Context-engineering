"""Mejora matemática y narrativa de la PPT de Context Engineering.

El archivo se edita con python-pptx. Las ecuaciones se renderizan como PNG
transparentes de alta resolución con MathText para conservar notación
matemática profesional dentro de PowerPoint.

Uso recomendado desde el entorno indicado por el proyecto:

    conda run -n ppt-env python outputs/mejorar_formulas_context_engineering.py

Si ``ppt-env`` no contiene Matplotlib, el script se relanza con el Python base
de Conda (que sí lo contiene) y reutiliza python-pptx desde ``ppt-env``.
"""

from __future__ import annotations

import argparse
import os
import re
import subprocess
import sys
import tempfile
from pathlib import Path


def _conda_base() -> Path | None:
    current = Path(sys.prefix).resolve()
    for candidate in (current, *current.parents):
        if candidate.name.lower() in {"anaconda3", "miniconda3", "miniforge3"}:
            return candidate
    fallback = Path.home() / "anaconda3"
    return fallback if fallback.exists() else None


def _bootstrap() -> None:
    """Hace compatible el comando solicitado con las librerías ya instaladas."""
    base = _conda_base()
    try:
        import matplotlib  # noqa: F401
    except ImportError:
        if "--_reexec" in sys.argv or base is None:
            raise RuntimeError("Matplotlib no está disponible en el entorno actual.")
        base_python = base / "python.exe"
        if not base_python.exists():
            raise RuntimeError(f"No se encontró el Python base de Conda: {base_python}")
        command = [str(base_python), str(Path(__file__).resolve()), *sys.argv[1:], "--_reexec"]
        raise SystemExit(subprocess.run(command, check=False).returncode)

    try:
        import pptx  # noqa: F401
    except ImportError:
        if base is None:
            raise
        candidates = [
            base / "envs" / "ppt-env" / "Lib" / "site-packages",
            base / "envs" / "ppt-env" / "lib" / "site-packages",
        ]
        for candidate in candidates:
            if candidate.exists():
                sys.path.append(str(candidate))
                break
        import pptx  # noqa: F401


_bootstrap()

import matplotlib as mpl
from matplotlib.backends.backend_agg import FigureCanvasAgg
from matplotlib.figure import Figure
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
DEFAULT_INPUT = ROOT / "outputs" / "Context_Engineering_y_Ambiguedad_refinada_python.pptx"
DEFAULT_OUTPUT = ROOT / "outputs" / "Context_Engineering_y_Ambiguedad_final_matematica.pptx"

PALETTE = {
    "ink": "101526",
    "white": "F7F8FC",
    "body_light": "60687C",
    "body_dark": "C7CEDF",
    "footer_light": "8E96AA",
    "footer_dark": "AAB3C8",
    "violet": "7065E8",
    "cyan": "3DD6C6",
    "orange": "F2A65A",
    "red": "F26674",
    "green": "52C97F",
    "line_light": "D9DEEA",
}


def normalized_text(value: str) -> str:
    return " ".join(value.split())


def iter_text_shapes(slide):
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False) and shape.text.strip():
            yield shape


def iter_runs(shape):
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.text:
                yield run


def find_text_shape(slide, exact_text: str):
    target = normalized_text(exact_text)
    for shape in iter_text_shapes(slide):
        if normalized_text(shape.text) == target:
            return shape
    raise ValueError(f"No se encontró el texto {exact_text!r}")


def remove_shape(shape) -> None:
    element = shape._element
    element.getparent().remove(element)


def set_text(shape, text: str, font_name="Aptos", size=14, color="60687C", bold=False,
             alignment=PP_ALIGN.LEFT) -> None:
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = alignment
    run = p.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def replace_plain_text(slide, old: str, new: str) -> None:
    shape = find_text_shape(slide, old)
    first_run = next(iter_runs(shape), None)
    font_name = first_run.font.name if first_run and first_run.font.name else "Aptos"
    font_size = first_run.font.size.pt if first_run and first_run.font.size else 14
    color = "101526"
    if first_run:
        try:
            if first_run.font.color.rgb:
                color = str(first_run.font.color.rgb)
        except (AttributeError, TypeError, ValueError):
            pass
    bold = bool(first_run.font.bold) if first_run else False
    alignment = shape.text_frame.paragraphs[0].alignment or PP_ALIGN.LEFT
    set_text(shape, new, font_name, font_size, color, bold, alignment)


def render_equation(path: Path, latex: str, color: str, fontsize: float = 34) -> None:
    """Renderiza una expresión MathText sobre fondo transparente."""
    mpl.rcParams.update({
        "mathtext.fontset": "stix",
        "font.family": "STIXGeneral",
        "savefig.transparent": True,
    })
    figure = Figure(figsize=(12, 1.7), dpi=320)
    FigureCanvasAgg(figure)
    figure.patch.set_alpha(0)
    figure.text(
        0.01,
        0.50,
        f"${latex}$",
        fontsize=fontsize,
        color=f"#{color}",
        ha="left",
        va="center",
    )
    figure.savefig(
        path,
        format="png",
        dpi=320,
        transparent=True,
        bbox_inches="tight",
        pad_inches=0.025,
    )


def add_picture_contain(slide, image_path: Path, left, top, width, height,
                        align="center", valign="center"):
    with Image.open(image_path) as image:
        img_w, img_h = image.size
    image_ratio = img_w / img_h
    box_ratio = width / height
    if image_ratio >= box_ratio:
        pic_w = width
        pic_h = int(width / image_ratio)
    else:
        pic_h = height
        pic_w = int(height * image_ratio)

    if align == "left":
        pic_left = left
    elif align == "right":
        pic_left = left + width - pic_w
    else:
        pic_left = left + (width - pic_w) // 2

    if valign == "top":
        pic_top = top
    elif valign == "bottom":
        pic_top = top + height - pic_h
    else:
        pic_top = top + (height - pic_h) // 2

    return slide.shapes.add_picture(str(image_path), pic_left, pic_top, pic_w, pic_h)


def add_equation(slide, tmp_dir: Path, key: str, latex: str, left, top, width, height,
                 color="101526", fontsize=34, align="center"):
    image_path = tmp_dir / f"{key}.png"
    render_equation(image_path, latex, color, fontsize)
    picture = add_picture_contain(slide, image_path, left, top, width, height, align=align)
    picture.name = f"Ecuación {key}"
    return picture


def replace_equation(slide, tmp_dir: Path, key: str, old_text: str, latex: str,
                     color="101526", fontsize=34, align="left", box=None):
    shape = find_text_shape(slide, old_text)
    if box is None:
        box = (shape.left, shape.top, shape.width, shape.height)
    remove_shape(shape)
    return add_equation(slide, tmp_dir, key, latex, *box, color=color,
                        fontsize=fontsize, align=align)


def add_legend(slide, text: str, left, top, width, height, dark=False, size=12.5):
    shape = slide.shapes.add_textbox(left, top, width, height)
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT

    label = p.add_run()
    label.text = "Notación · "
    label.font.name = "Aptos"
    label.font.size = Pt(size)
    label.font.bold = True
    label.font.color.rgb = RGBColor.from_string(PALETTE["cyan"] if dark else PALETTE["violet"])

    body = p.add_run()
    body.text = text
    body.font.name = "Aptos"
    body.font.size = Pt(size)
    body.font.color.rgb = RGBColor.from_string(PALETTE["body_dark"] if dark else PALETTE["body_light"])
    return shape


def add_badge(slide, text: str, left, top, width, height, color="7065E8"):
    badge = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height)
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor.from_string("EEF0F8")
    badge.line.color.rgb = RGBColor.from_string(color)
    badge.line.width = Pt(1)
    set_text(badge, text, "Aptos", 11, color, True, PP_ALIGN.CENTER)
    return badge


def style_card(shape, bg_color=None, border_color=None, border_width=1, rounded=True) -> None:
    try:
        if bg_color:
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string(bg_color)
        else:
            shape.fill.background()
            
        if border_color:
            shape.line.color.rgb = RGBColor.from_string(border_color)
            shape.line.width = Pt(border_width)
        else:
            shape.line.fill.background()
    except Exception:
        pass



def update_equations(prs: Presentation, tmp_dir: Path) -> None:
    # 04 · Densidad de señal.
    slide = prs.slides[3]
    replace_equation(
        slide, tmp_dir, "s04_signal_density", "I(Y; C | X) / |C|",
        r"\frac{I(Y;C\mid X)}{|C|}", color=PALETTE["white"], fontsize=42,
    )
    add_legend(slide, "X consulta · Y respuesta · C contexto · |C| número de tokens",
               Inches(0.75), Inches(4.45), Inches(5.9), Inches(0.35), dark=True)

    # 05 · Distribución autoregresiva.
    slide = prs.slides[4]
    replace_equation(slide, tmp_dir, "s05_token", "Pθ(yₜ | X, C, y<ₜ)",
                     r"p_\theta(y_t\mid X,C,y_{<t})", color=PALETTE["violet"], fontsize=41)
    replace_equation(slide, tmp_dir, "s05_sequence",
                     "Pθ(Y | X, C) = ∏ₜ Pθ(yₜ | X, C, y<ₜ)",
                     r"p_\theta(Y\mid X,C)=\prod_{t=1}^{T}p_\theta(y_t\mid X,C,y_{<t})",
                     fontsize=32)
    replace_equation(slide, tmp_dir, "s05_example",
                     "P(“24” | X,C) > P(“12” | X,C)",
                     r"p(y_t=24\mid X,C)>p(y_t=12\mid X,C)", fontsize=25, align="center")
    add_legend(slide, "θ parámetros del LLM · t posición · yₜ token actual · y<ₜ prefijo generado",
               Inches(0.75), Inches(6.03), Inches(10.8), Inches(0.32))

    # 06 · Entropía e información mutua.
    slide = prs.slides[5]
    replace_equation(slide, tmp_dir, "s06_entropy", "H(Y | X, C) ≤ H(Y | X)",
                     r"H(Y\mid X,C)\leq H(Y\mid X)", fontsize=39)
    replace_equation(slide, tmp_dir, "s06_mutual",
                     "I(Y; C | X) = H(Y | X) − H(Y | X, C)",
                     r"I(Y;C\mid X)=H(Y\mid X)-H(Y\mid X,C)",
                     color=PALETTE["violet"], fontsize=28,
                     box=(Inches(0.75), Inches(3.15), Inches(6.15), Inches(0.55)))
    add_legend(slide,
               "H entropía · I información mutua · X consulta · Y respuesta · C contexto\nVálido para un predictor ideal, en promedio sobre C.",
               Inches(0.75), Inches(4.55), Inches(6.15), Inches(0.68), size=11.5)

    # 07 · Modelo conceptual, no identidad probabilística.
    slide = prs.slides[6]
    replace_equation(slide, tmp_dir, "s07_quality",
                     "Q(C) ≈ R(C) − N(C) − D(C) − K(C)",
                     r"Q(C)\approx R(C)-N(C)-D(C)-K(C)", fontsize=35)
    replace_equation(slide, tmp_dir, "s07_nonmonotonic", "Q(n + 1) ≱ Q(n)",
                     r"Q(n+1)\ngeq Q(n)", color=PALETTE["violet"], fontsize=31)
    add_badge(slide, "MODELO CONCEPTUAL", Inches(0.75), Inches(2.08), Inches(1.95), Inches(0.28))
    add_legend(slide, "Q calidad observada · R relevancia · N ruido · D dificultad · K conflictos · n tokens",
               Inches(0.75), Inches(5.15), Inches(6.1), Inches(0.48))

    # 08 · Descomposición end-to-end de RAG.
    slide = prs.slides[7]
    replace_equation(
        slide, tmp_dir, "s08_rag_product",
        "P(correcta) ≈ P(recuperar evidencia) × P(responder | evidencia)",
        r"p_{\mathrm{RAG}}(Y^\star\mid X,D)\approx p_\eta(C^\star\mid X,D)\,p_\theta(Y^\star\mid X,C^\star)",
        fontsize=30, align="center",
    )
    add_legend(slide, "D corpus · C★ evidencia correcta · Y★ respuesta correcta · η retriever · θ generador",
               Inches(1.1), Inches(3.10), Inches(11.0), Inches(0.32))

    # 09 actual · Atención; pasará a ser la diapositiva 10 tras insertar la comparación.
    slide = prs.slides[8]
    replace_equation(slide, tmp_dir, "s09_attention",
                     "Attention(Q,K,V) = softmax(QKᵀ / √dₖ)V",
                     r"\mathrm{Attention}(Q,K,V)=\mathrm{softmax}\!\left(\frac{QK^\top}{\sqrt{d_k}}\right)V",
                     color=PALETTE["white"], fontsize=31)
    add_legend(slide, "Q consultas · K claves · V valores · dₖ dimensión de las claves",
               Inches(0.75), Inches(3.62), Inches(6.35), Inches(0.38), dark=True)

    # 11 actual · Mezcla sobre intención latente.
    slide = prs.slides[10]
    replace_equation(slide, tmp_dir, "s11_mixture",
                     "P(Y | X,C) = Σz P(Y | X,C,Z=z) · P(Z=z | X,C)",
                     r"p(Y\mid X,C)=\sum_z p(Y\mid X,C,Z=z)\,p(Z=z\mid X,C)",
                     color=PALETTE["violet"], fontsize=33, align="center")
    add_legend(slide, "Z intención latente · z hipótesis de intención · la salida marginaliza interpretaciones posibles",
               Inches(1.1), Inches(3.13), Inches(11.0), Inches(0.34))

    # 12 actual · Sustituye una implicación no universal por la regla exacta de la cadena.
    slide = prs.slides[11]
    replace_plain_text(slide, "La incertidumbre sobre Z se propaga hasta Y",
                       "La incertidumbre combina intención y respuesta")
    for text in ("H(Z | X) ↑", "⇒", "H(Y | X,C) ↑", "riesgo ↑"):
        shape = find_text_shape(slide, text)
        remove_shape(shape)
    # Hay dos flechas iguales; se elimina la segunda por separado.
    remaining_arrows = [shape for shape in iter_text_shapes(slide) if normalized_text(shape.text) == "⇒"]
    for shape in remaining_arrows:
        remove_shape(shape)
    add_equation(slide, tmp_dir, "s12_chain_rule",
                 r"H(Y,Z\mid X,C)=H(Z\mid X,C)+H(Y\mid X,C,Z)",
                 Inches(0.75), Inches(2.28), Inches(11.8), Inches(0.62),
                 color=PALETTE["white"], fontsize=35, align="center")
    note = slide.shapes.add_textbox(Inches(1.75), Inches(2.96), Inches(9.85), Inches(0.35))
    set_text(note, "incertidumbre total  =  intención no resuelta  +  respuesta condicionada a Z",
             "Aptos", 16, PALETTE["cyan"], True, PP_ALIGN.CENTER)
    add_legend(slide, "H entropía · X consulta · C contexto · Z intención · Y respuesta",
               Inches(1.75), Inches(3.31), Inches(9.85), Inches(0.32), dark=True)
    replace_plain_text(slide, "Fuente: context.md — entropía de intención y multimodalidad",
                       "Fuente: context.md — regla de la cadena de entropía y multimodalidad")

    # 13 actual · Embedding ambiguo y margen de recuperación.
    slide = prs.slides[12]
    replace_equation(slide, tmp_dir, "s13_embedding",
                     "eX ≈ Σz P(Z=z | X) · eX,z",
                     r"e_X\approx\sum_z p(Z=z\mid X)\,e_{X,z}",
                     color=PALETTE["violet"], fontsize=29)
    replace_equation(slide, tmp_dir, "s13_margin",
                     "Δ = s(Ccorrecto,X) − s(Cincorrecto,X) ≈ 0",
                     r"\Delta=s(C^+,X)-s(C^-,X)\approx 0", fontsize=29)
    add_badge(slide, "APROX. CONCEPTUAL", Inches(0.75), Inches(1.68), Inches(1.58), Inches(0.27))
    add_legend(slide, "eₓ embedding · s similitud · C⁺ candidato relevante · C⁻ distractor · Δ margen",
               Inches(0.75), Inches(4.73), Inches(5.65), Inches(0.52))
    replace_plain_text(slide,
                       "Fuente: context.md — embedding agregado y reducción del margen de similitud",
                       "Fuente: context.md — aproximación conceptual del embedding y margen de similitud")

    # 14 actual · Reranking marginalizado por intención.
    slide = prs.slides[13]
    priors = find_text_shape(slide, "P(Z₁|X)=0.55 · P(Z₂|X)=0.45")
    remove_shape(priors)
    add_equation(slide, tmp_dir, "s14_rerank",
                 r"s_i=\sum_z p(C_i\,\mathrm{rel}\mid X,Z=z)\,p(Z=z\mid X)",
                 Inches(1.05), Inches(1.93), Inches(11.2), Inches(0.43),
                 fontsize=27, align="center")
    add_equation(slide, tmp_dir, "s14_priors",
                 r"p(Z_1\mid X)=0.55\qquad p(Z_2\mid X)=0.45",
                 Inches(1.05), Inches(2.40), Inches(11.2), Inches(0.35),
                 color=PALETTE["body_light"], fontsize=23, align="center")
    replace_equation(slide, tmp_dir, "s14_delta", "Δ = 0.0025",
                     r"\Delta=0.0025", color=PALETTE["red"], fontsize=25, align="center")
    add_legend(slide, "sᵢ score marginal · Cᵢ chunk · Z intención latente · p(Z=z|X) peso de cada interpretación",
               Inches(1.15), Inches(6.08), Inches(10.9), Inches(0.36))

    # 16 actual · Riesgo Bayesiano bajo pérdida 0–1.
    slide = prs.slides[15]
    replace_equation(slide, tmp_dir, "s16_prior1", "P(Z₁|X)=0.60",
                     r"p(Z_1\mid X)=0.60", color=PALETTE["cyan"], fontsize=29)
    replace_equation(slide, tmp_dir, "s16_prior2", "P(Z₂|X)=0.40",
                     r"p(Z_2\mid X)=0.40", color=PALETTE["violet"], fontsize=29)
    replace_equation(slide, tmp_dir, "s16_bayes", "R* = 1 − maxz P(Z=z|X)",
                     r"R^\star(X)=1-\max_z p(Z=z\mid X)",
                     color=PALETTE["white"], fontsize=31, align="center")
    replace_equation(slide, tmp_dir, "s16_result", "= 40%", r"=40\%",
                     color=PALETTE["red"], fontsize=42, align="center")
    add_legend(slide, "R★ riesgo Bayesiano mínimo · Z intención · pérdida 0–1",
               Inches(5.75), Inches(4.13), Inches(5.85), Inches(0.32), dark=True, size=11)

    # 17 actual · Información de una pregunta aclaratoria.
    slide = prs.slides[16]
    replace_equation(slide, tmp_dir, "s17_information",
                     "I(Z;A|X) = H(Z|X) − H(Z|X,A)",
                     r"I(Z;A_q\mid X)=H(Z\mid X)-H(Z\mid X,A_q)",
                     color=PALETTE["violet"], fontsize=29, align="center")
    replace_equation(slide, tmp_dir, "s17_question",
                     "A* = arg maxA I(Z;A|X)",
                     r"q^\star=\arg\max_q I(Z;A_q\mid X)", fontsize=29, align="center")
    add_legend(slide, "q pregunta · A_q respuesta · Z intención · I información ganada",
               Inches(7.25), Inches(3.92), Inches(4.85), Inches(0.38), size=11.5)

    # 18 actual · Decisión Bayesiana con costo de interacción.
    slide = prs.slides[17]
    condition = find_text_shape(slide, "Raclarar < Rdirecto")
    remove_shape(condition)
    add_equation(slide, tmp_dir, "s18_direct",
                 r"R_{\mathrm{dir}}(X)=\min_{\hat Y}\;\mathbb{E}[L(\hat Y,Y)\mid X,C]",
                 Inches(0.85), Inches(1.92), Inches(5.45), Inches(0.44),
                 color=PALETTE["white"], fontsize=23, align="center")
    add_equation(slide, tmp_dir, "s18_clarify",
                 r"R_{\mathrm{acl}}(X,q)=c_q+\mathbb{E}_{A_q}\!\left[\min_{\hat Y}\mathbb{E}[L(\hat Y,Y)\mid X,C,A_q]\right]",
                 Inches(6.45), Inches(1.92), Inches(6.05), Inches(0.44),
                 color=PALETTE["white"], fontsize=20, align="center")
    add_equation(slide, tmp_dir, "s18_condition",
                 r"\min_q R_{\mathrm{acl}}(X,q)<R_{\mathrm{dir}}(X)",
                 Inches(3.35), Inches(2.48), Inches(6.65), Inches(0.48),
                 color=PALETTE["cyan"], fontsize=30, align="center")
    variables = find_text_shape(slide,
                                "{Costo de error} · {Costo de interacción} · {Umbral de ambigüedad}")
    set_text(variables,
             "L: pérdida · c_q: costo de preguntar · q: pregunta · A_q: respuesta · Ŷ: decisión",
             "Aptos", 13, PALETTE["body_dark"], False, PP_ALIGN.CENTER)

    # 20 actual · La métrica de detección también se compone como ecuación.
    slide = prs.slides[19]
    replace_equation(slide, tmp_dir, "s20_entropy", "H(Z|X)", r"H(Z\mid X)",
                     color=PALETTE["body_dark"], fontsize=20, align="center")
    add_legend(slide, "H entropía · Z intención · X consulta",
               Inches(0.75), Inches(6.32), Inches(5.0), Inches(0.34), dark=True)


def add_new_comparison_slide(prs: Presentation, tmp_dir: Path):
    # Algunas PPTX generadas desde cero solo conservan uno o dos layouts.
    # Se elige el layout con menos placeholders y se limpia antes de componer.
    layout = min(prs.slide_layouts, key=lambda item: len(item.placeholders))
    slide = prs.slides.add_slide(layout)
    for placeholder in list(slide.placeholders):
        remove_shape(placeholder)
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = RGBColor.from_string(PALETTE["white"])

    kicker = slide.shapes.add_textbox(Inches(0.75), Inches(0.40), Inches(4.5), Inches(0.24))
    set_text(kicker, "09 / MODELOS COMPARADOS", "Aptos", 12, PALETTE["violet"], True)

    title = slide.shapes.add_textbox(Inches(0.75), Inches(0.76), Inches(11.85), Inches(0.64))
    set_text(title, "RAG recupera; Context Engineering orquesta", "Aptos Display", 36,
             PALETTE["ink"], True)

    divider = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE,
                                     Inches(0.75), Inches(1.60), Inches(11.83), Inches(0.012))
    divider.fill.solid()
    divider.fill.fore_color.rgb = RGBColor.from_string(PALETTE["line_light"])
    divider.line.fill.background()

    # RAG Card background
    rag_card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.65), Inches(1.75), Inches(5.70), Inches(3.15))
    style_card(rag_card, bg_color="F1F3F9", border_color=PALETTE["line_light"], border_width=1, rounded=True)
    slide.shapes._spTree.remove(rag_card._element)
    slide.shapes._spTree.insert(2, rag_card._element)

    # CE Card background
    ce_card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(6.75), Inches(1.75), Inches(5.70), Inches(3.15))
    style_card(ce_card, bg_color="F0EEFF", border_color=PALETTE["violet"], border_width=1.5, rounded=True)
    slide.shapes._spTree.remove(ce_card._element)
    slide.shapes._spTree.insert(3, ce_card._element)

    left_label = slide.shapes.add_textbox(Inches(0.78), Inches(1.88), Inches(1.1), Inches(0.28))
    set_text(left_label, "RAG", "Aptos", 16, PALETTE["violet"], True)
    left_claim = slide.shapes.add_textbox(Inches(0.78), Inches(2.20), Inches(5.45), Inches(0.48))
    set_text(left_claim, "Recupera C desde documentos D", "Aptos Display", 20,
             PALETTE["ink"], True)
    add_equation(slide, tmp_dir, "s09_rag_model",
                 r"p_{\mathrm{RAG}}(Y\mid X,D)=\sum_{C\in\mathcal{C}(D)}p_\theta(Y\mid X,C)\,p_\eta(C\mid X,D)",
                 Inches(0.78), Inches(2.85), Inches(5.47), Inches(0.92),
                 color=PALETTE["ink"], fontsize=22, align="center")
    left_note = slide.shapes.add_textbox(Inches(0.78), Inches(3.92), Inches(5.45), Inches(0.72))
    set_text(left_note, "η recupera documentos; θ genera con el contexto seleccionado.",
             "Aptos", 15, PALETTE["body_light"], False)

    right_label = slide.shapes.add_textbox(Inches(6.98), Inches(1.88), Inches(3.2), Inches(0.28))
    set_text(right_label, "CONTEXT ENGINEERING", "Aptos", 16, PALETTE["cyan"], True)
    right_claim = slide.shapes.add_textbox(Inches(6.98), Inches(2.20), Inches(5.35), Inches(0.48))
    set_text(right_claim, "Orquesta C desde el estado S", "Aptos Display", 20,
             PALETTE["ink"], True)
    add_equation(slide, tmp_dir, "s09_ce_model",
                 r"p_{\mathrm{CE}}(Y\mid X,S)=\sum_{C\in\mathcal{C}(S)}p_\theta(Y\mid X,C)\,p_\phi(C\mid X,S)",
                 Inches(6.98), Inches(2.85), Inches(5.35), Inches(0.92),
                 color=PALETTE["ink"], fontsize=22, align="center")
    add_equation(slide, tmp_dir, "s09_state",
                 r"S=\{D,M,T,H,E\}",
                 Inches(7.75), Inches(3.84), Inches(3.9), Inches(0.43),
                 color=PALETTE["cyan"], fontsize=25, align="center")
    right_note = slide.shapes.add_textbox(Inches(6.98), Inches(4.30), Inches(5.35), Inches(0.50))
    set_text(right_note, "φ combina documentos, memoria, herramientas, historial y entorno.",
             "Aptos", 15, PALETTE["body_light"], False)

    takeaway = slide.shapes.add_textbox(Inches(1.15), Inches(5.08), Inches(7.3), Inches(0.50))
    set_text(takeaway, "RAG es un caso particular", "Aptos Display", 24,
             PALETTE["violet"], True, PP_ALIGN.CENTER)
    add_equation(slide, tmp_dir, "s09_special_case", r"S=D,\quad\phi=\eta",
                 Inches(8.20), Inches(5.08), Inches(3.1), Inches(0.48),
                 color=PALETTE["violet"], fontsize=27, align="center")

    add_legend(slide,
               "X consulta · Y respuesta · C contexto · D documentos · M memoria · T herramientas\nH historial · E entorno · η retriever · φ orquestador",
               Inches(0.78), Inches(5.78), Inches(11.55), Inches(0.72), size=11.5)

    source = slide.shapes.add_textbox(Inches(0.75), Inches(7.04), Inches(10.9), Inches(0.20))
    set_text(source, "Fuente: formulación matemática del proyecto — RAG como caso particular de Context Engineering",
             "Aptos", 10.5, PALETTE["footer_light"], False)
    page = slide.shapes.add_textbox(Inches(12.08), Inches(7.02), Inches(0.50), Inches(0.22))
    set_text(page, "09", "Aptos", 11, PALETTE["footer_light"], True, PP_ALIGN.RIGHT)

    # Mueve el slide recién creado (último) después del RAG end-to-end (posición 9).
    slide_id_list = prs.slides._sldIdLst
    new_id = slide_id_list[-1]
    slide_id_list.remove(new_id)
    slide_id_list.insert(8, new_id)


def add_conversational_challenges_slide(prs: Presentation):
    layout = min(prs.slide_layouts, key=lambda item: len(item.placeholders))
    slide = prs.slides.add_slide(layout)
    for placeholder in list(slide.placeholders):
        remove_shape(placeholder)
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = RGBColor.from_string(PALETTE["white"])

    kicker = slide.shapes.add_textbox(Inches(0.75), Inches(0.40), Inches(5.5), Inches(0.24))
    set_text(kicker, "11 / DESAFÍOS CONVERSIONALES", "Aptos", 12, PALETTE["violet"], True)

    title = slide.shapes.add_textbox(Inches(0.75), Inches(0.76), Inches(11.85), Inches(0.64))
    set_text(title, "Cuando X no representa una intención literal", "Aptos Display", 36,
             PALETTE["ink"], True)

    divider = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE,
                                     Inches(0.75), Inches(1.60), Inches(11.83), Inches(0.012))
    divider.fill.solid()
    divider.fill.fore_color.rgb = RGBColor.from_string(PALETTE["line_light"])
    divider.line.fill.background()

    # 4 Cards coordinates:
    # 1. Ironía (top left)
    card1 = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.65), Inches(1.75), Inches(5.70), Inches(1.45))
    style_card(card1, bg_color="FDF0F0", border_color=PALETTE["red"], border_width=1, rounded=True)
    slide.shapes._spTree.remove(card1._element)
    slide.shapes._spTree.insert(2, card1._element)

    t1 = slide.shapes.add_textbox(Inches(0.78), Inches(1.85), Inches(5.45), Inches(1.25))
    tf1 = t1.text_frame
    tf1.clear()
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    r = p.add_run()
    r.text = "1. Ironía y Sarcasmo\n"
    r.font.name = "Aptos"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(PALETTE["red"])
    
    r2 = p.add_run()
    r2.text = "Inversión de la polaridad semántica literal.\n"
    r2.font.name = "Aptos"
    r2.font.size = Pt(12)
    r2.font.color.rgb = RGBColor.from_string(PALETTE["body_light"])
    
    r3 = p.add_run()
    r3.text = "X: \"Qué gran servicio...\" (esperando 3 semanas) → Z: Reclamo"
    r3.font.name = "Consolas"
    r3.font.size = Pt(11.5)
    r3.font.color.rgb = RGBColor.from_string(PALETTE["ink"])
    r3.font.bold = True

    # 2. Actos de Habla Indirectos (top right)
    card2 = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(6.75), Inches(1.75), Inches(5.70), Inches(1.45))
    style_card(card2, bg_color="EEF0F8", border_color=PALETTE["violet"], border_width=1, rounded=True)
    slide.shapes._spTree.remove(card2._element)
    slide.shapes._spTree.insert(3, card2._element)

    t2 = slide.shapes.add_textbox(Inches(6.88), Inches(1.85), Inches(5.45), Inches(1.25))
    tf2 = t2.text_frame
    tf2.clear()
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    r = p.add_run()
    r.text = "2. Actos de Habla Indirectos\n"
    r.font.name = "Aptos"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(PALETTE["violet"])
    
    r2 = p.add_run()
    r2.text = "Peticiones de acción formuladas como preguntas de habilidad.\n"
    r2.font.name = "Aptos"
    r2.font.size = Pt(12)
    r2.font.color.rgb = RGBColor.from_string(PALETTE["body_light"])
    
    r3 = p.add_run()
    r3.text = "X: \"¿Tienes el saldo disponible?\" → Z: Comando de consulta"
    r3.font.name = "Consolas"
    r3.font.size = Pt(11.5)
    r3.font.color.rgb = RGBColor.from_string(PALETTE["ink"])
    r3.font.bold = True

    # 3. Elipsis y Co-referencia (bottom left)
    card3 = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.65), Inches(3.35), Inches(5.70), Inches(1.45))
    style_card(card3, bg_color="E8FCFA", border_color=PALETTE["cyan"], border_width=1, rounded=True)
    slide.shapes._spTree.remove(card3._element)
    slide.shapes._spTree.insert(4, card3._element)

    t3 = slide.shapes.add_textbox(Inches(0.78), Inches(3.45), Inches(5.45), Inches(1.25))
    tf3 = t3.text_frame
    tf3.clear()
    tf3.word_wrap = True
    p = tf3.paragraphs[0]
    r = p.add_run()
    r.text = "3. Elipsis y Co-referencia\n"
    r.font.name = "Aptos"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(PALETTE["cyan"])
    
    r2 = p.add_run()
    r2.text = "Omisión de información que se asume conocida por el contexto.\n"
    r2.font.name = "Aptos"
    r2.font.size = Pt(12)
    r2.font.color.rgb = RGBColor.from_string(PALETTE["body_light"])
    
    r3 = p.add_run()
    r3.text = "X: \"Haz lo mismo con la otra\" → Z: Aplicar acción al historial"
    r3.font.name = "Consolas"
    r3.font.size = Pt(11.5)
    r3.font.color.rgb = RGBColor.from_string(PALETTE["ink"])
    r3.font.bold = True

    # 4. Modismos y Lenguaje Figurado (bottom right)
    card4 = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(6.75), Inches(3.35), Inches(5.70), Inches(1.45))
    style_card(card4, bg_color="FFF7EE", border_color=PALETTE["orange"], border_width=1, rounded=True)
    slide.shapes._spTree.remove(card4._element)
    slide.shapes._spTree.insert(5, card4._element)

    t4 = slide.shapes.add_textbox(Inches(6.88), Inches(3.45), Inches(5.45), Inches(1.25))
    tf4 = t4.text_frame
    tf4.clear()
    tf4.word_wrap = True
    p = tf4.paragraphs[0]
    r = p.add_run()
    r.text = "4. Modismos y Lenguaje Figurado\n"
    r.font.name = "Aptos"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(PALETTE["orange"])
    
    r2 = p.add_run()
    r2.text = "Semántica no composicional y modismos locales.\n"
    r2.font.name = "Aptos"
    r2.font.size = Pt(12)
    r2.font.color.rgb = RGBColor.from_string(PALETTE["body_light"])
    
    r3 = p.add_run()
    r3.text = "X: \"Se me cayó el sistema\" → Z: Reporte de fallo técnico"
    r3.font.name = "Consolas"
    r3.font.size = Pt(11.5)
    r3.font.color.rgb = RGBColor.from_string(PALETTE["ink"])
    r3.font.bold = True

    # Takeaway at the bottom
    takeaway = slide.shapes.add_textbox(Inches(1.15), Inches(5.00), Inches(11.0), Inches(0.60))
    set_text(takeaway, "Al igual que la ambigüedad matemática, estos desafíos requieren que el sistema infiera una intención latente Z a partir de una consulta X que no describe literalmente la acción esperada.",
             "Aptos Display", 18, PALETTE["violet"], True, PP_ALIGN.CENTER)

    add_legend(slide,
               "X consulta observable del usuario · Z intención latente real que el agente debe resolver para actuar correctamente",
               Inches(0.78), Inches(5.78), Inches(11.55), Inches(0.72), size=11.5)

    source = slide.shapes.add_textbox(Inches(0.75), Inches(7.04), Inches(10.9), Inches(0.20))
    set_text(source, "Fuente: context.md — pragmática del lenguaje natural y teoría de actos de habla",
             "Aptos", 10.5, PALETTE["footer_light"], False)
    page = slide.shapes.add_textbox(Inches(12.08), Inches(7.02), Inches(0.50), Inches(0.22))
    set_text(page, "11", "Aptos", 11, PALETTE["footer_light"], True, PP_ALIGN.RIGHT)

    # Move slide to index 11
    slide_id_list = prs.slides._sldIdLst
    new_id = slide_id_list[-1]
    slide_id_list.remove(new_id)
    slide_id_list.insert(11, new_id)


def renumber_slides(prs: Presentation) -> None:
    for number, slide in enumerate(prs.slides, 1):
        # Número inferior derecho.
        for shape in iter_text_shapes(slide):
            if shape.top > Inches(6.75) and re.fullmatch(r"\d{2}", shape.text.strip()):
                first = next(iter_runs(shape), None)
                color = PALETTE["footer_light"]
                if first:
                    try:
                        if first.font.color.rgb:
                            color = str(first.font.color.rgb)
                    except (AttributeError, TypeError, ValueError):
                        pass
                set_text(shape, f"{number:02d}", "Aptos", 11, color, True, PP_ALIGN.RIGHT)

        if number == 1:
            continue
        # Kicker superior. La portada de la parte II conserva su tratamiento editorial.
        candidates = [shape for shape in iter_text_shapes(slide) if shape.top < Inches(0.72)]
        if not candidates:
            continue
        kicker = min(candidates, key=lambda shape: (shape.top, shape.left))
        current = normalized_text(kicker.text)
        if current.startswith("PARTE II"):
            continue
        current = re.sub(r"^\d{2}\s*/\s*", "", current)
        first = next(iter_runs(kicker), None)
        color = PALETTE["violet"]
        if first:
            try:
                if first.font.color.rgb:
                    color = str(first.font.color.rgb)
            except (AttributeError, TypeError, ValueError):
                pass
        set_text(kicker, f"{number:02d} / {current}", "Aptos", 12, color, True)


def apply_styling_improvements(prs: Presentation) -> None:
    # index 0: Portada
    slide = prs.slides[0]
    try:
        for shape in slide.shapes:
            if shape.has_text_frame and "{Nombre}" in shape.text:
                set_text(shape, "Kevin Gomez Villanueva  ·  AI Engineer en BCP  ·  Julio 2026", 
                         font_name="Aptos", size=13, color=PALETTE["body_dark"], alignment=PP_ALIGN.LEFT)
    except Exception:
        pass

    # index 1: El Problema
    slide = prs.slides[1]
    for shape in slide.shapes:
        if shape.has_text_frame:
            y_pt = shape.top / 12700
            text = shape.text.strip().lower()
            if abs(y_pt - 355) < 10 and text in {"activos", "capitalización", "clientes", "ingresos", "presencia"}:
                style_card(shape, bg_color="EEF0F8", border_color=PALETTE["violet"], border_width=1, rounded=True)
                set_text(shape, shape.text.strip(), font_name="Aptos", size=13, color=PALETTE["violet"], bold=True, alignment=PP_ALIGN.CENTER)

    # index 2: Framework
    slide = prs.slides[2]
    for shape in slide.shapes:
        x_pt = shape.left / 12700
        y_pt = shape.top / 12700
        if abs(x_pt - 530) < 15 and abs(y_pt - 184) < 15:
            style_card(shape, bg_color="F1F3F9", border_color=PALETTE["line_light"], border_width=1, rounded=True)

    # index 3: Tesis Central
    slide = prs.slides[3]
    for shape in slide.shapes:
        x_pt = shape.left / 12700
        y_pt = shape.top / 12700
        w_pt = shape.width / 12700
        h_pt = shape.height / 12700
        if abs(x_pt - 760) < 10 and abs(y_pt - 231) < 10:
            if abs(w_pt - 420) < 10:
                style_card(shape, bg_color="FDF0F0", border_color=PALETTE["red"], border_width=1, rounded=True)
            elif abs(w_pt - 105) < 10:
                style_card(shape, bg_color=PALETTE["violet"], border_color=PALETTE["violet"], border_width=1, rounded=True)
        elif abs(y_pt - 454) < 10 and shape.has_text_frame:
            text = shape.text.strip().lower()
            if text in {"seleccionar", "ordenar", "comprimir"}:
                style_card(shape, bg_color="EEF0F8", border_color=PALETTE["violet"], border_width=1, rounded=True)
                set_text(shape, shape.text.strip(), font_name="Aptos", size=13, color=PALETTE["violet"], bold=True, alignment=PP_ALIGN.CENTER)

    # index 4: Modelo Probabilístico
    slide = prs.slides[4]
    for shape in slide.shapes:
        x_pt = shape.left / 12700
        y_pt = shape.top / 12700
        if abs(x_pt - 700) < 10 and abs(y_pt - 242) < 10:
            style_card(shape, bg_color="F1F3F9", border_color=PALETTE["line_light"], border_width=1, rounded=True)

    # index 5: Valor de la Información
    slide = prs.slides[5]
    for shape in slide.shapes:
        x_pt = shape.left / 12700
        y_pt = shape.top / 12700
        w_pt = shape.width / 12700
        if abs(x_pt - 780) < 10:
            if abs(y_pt - 238) < 10:
                if abs(w_pt - 340) < 10:
                    style_card(shape, bg_color="EAECEF", border_color=None, rounded=True)
                elif abs(w_pt - 305) < 10:
                    style_card(shape, bg_color=PALETTE["red"], border_color=None, rounded=True)
            elif abs(y_pt - 366) < 10:
                if abs(w_pt - 340) < 10:
                    style_card(shape, bg_color="EAECEF", border_color=None, rounded=True)
                elif abs(w_pt - 112) < 10:
                    style_card(shape, bg_color=PALETTE["violet"], border_color=None, rounded=True)

    # index 6: Límites del Modelo
    slide = prs.slides[6]
    for shape in slide.shapes:
        x_pt = shape.left / 12700
        y_pt = shape.top / 12700
        w_pt = shape.width / 12700
        h_pt = shape.height / 12700
        if abs(y_pt - 339) < 10 and abs(w_pt - 52) < 5 and abs(h_pt - 52) < 5:
            text = shape.text.strip().upper()
            if text == "R":
                style_card(shape, bg_color="ECFDF3", border_color=PALETTE["green"], border_width=1.5, rounded=True)
                set_text(shape, "R", font_name="Aptos", size=18, color=PALETTE["green"], bold=True, alignment=PP_ALIGN.CENTER)
            elif text == "N":
                style_card(shape, bg_color="FDF0F0", border_color=PALETTE["red"], border_width=1.5, rounded=True)
                set_text(shape, "N", font_name="Aptos", size=18, color=PALETTE["red"], bold=True, alignment=PP_ALIGN.CENTER)
            elif text == "D":
                style_card(shape, bg_color="FFF7EE", border_color=PALETTE["orange"], border_width=1.5, rounded=True)
                set_text(shape, "D", font_name="Aptos", size=18, color=PALETTE["orange"], bold=True, alignment=PP_ALIGN.CENTER)
            elif text == "K":
                style_card(shape, bg_color="FDF0F0", border_color=PALETTE["red"], border_width=1.5, rounded=True)
                set_text(shape, "K", font_name="Aptos", size=18, color=PALETTE["red"], bold=True, alignment=PP_ALIGN.CENTER)
        elif abs(w_pt - 60) < 5 or abs(w_pt - 40) < 5:
            if x_pt >= 750 and x_pt <= 1130 and y_pt >= 300 and y_pt <= 500:
                if x_pt < 900:
                    style_card(shape, bg_color="7065E8", border_color=None, rounded=False)
                elif x_pt < 980:
                    style_card(shape, bg_color="3DD6C6", border_color=None, rounded=False)
                else:
                    style_card(shape, bg_color="F26674", border_color=None, rounded=False)
        elif abs(x_pt - 922) < 5 and abs(y_pt - 322) < 5:
            style_card(shape, bg_color=PALETTE["violet"], border_color=PALETTE["white"], border_width=1.5, rounded=True)

    # index 7: RAG End-to-End
    slide = prs.slides[7]
    for shape in slide.shapes:
        x_pt = shape.left / 12700
        y_pt = shape.top / 12700
        w_pt = shape.width / 12700
        h_pt = shape.height / 12700
        if abs(y_pt - 354) < 10 and abs(h_pt - 140) < 10:
            if abs(x_pt - 130) < 10:
                style_card(shape, bg_color="F1F3F9", border_color=PALETTE["line_light"], border_width=1, rounded=True)
            elif abs(x_pt - 520) < 10:
                style_card(shape, bg_color="F1F3F9", border_color=PALETTE["line_light"], border_width=1, rounded=True)
            elif abs(x_pt - 930) < 10:
                style_card(shape, bg_color="F0EEFF", border_color=PALETTE["violet"], border_width=1.5, rounded=True)
        elif shape.has_text_frame:
            text = shape.text.strip()
            if text == "0.72":
                set_text(shape, "0.72", font_name="Aptos Display", size=36, color=PALETTE["violet"], bold=True, alignment=PP_ALIGN.CENTER)
            elif text in {"0.80", "0.90"}:
                set_text(shape, text, font_name="Aptos Display", size=36, color=PALETTE["ink"], bold=True, alignment=PP_ALIGN.CENTER)
            elif text == "SISTEMA":
                set_text(shape, "SISTEMA", font_name="Aptos", size=12, color=PALETTE["violet"], bold=True, alignment=PP_ALIGN.CENTER)
            elif text in {"RETRIEVER", "GENERATOR"}:
                set_text(shape, text, font_name="Aptos", size=12, color=PALETTE["body_light"], bold=True, alignment=PP_ALIGN.CENTER)

    # index 9: Attention (formerly 8)
    slide = prs.slides[9]
    for shape in slide.shapes:
        x_pt = shape.left / 12700
        y_pt = shape.top / 12700
        w_pt = shape.width / 12700
        h_pt = shape.height / 12700
        if abs(w_pt - 34) < 2 and abs(h_pt - 34) < 2:
            if x_pt >= 820 and x_pt <= 960:
                style_card(shape, bg_color=PALETTE["violet"], border_color=None, rounded=True)
            elif x_pt >= 960 and x_pt <= 1060:
                style_card(shape, bg_color=PALETTE["cyan"], border_color=None, rounded=True)
            else:
                style_card(shape, bg_color="EAECEF", border_color=None, rounded=True)
        elif abs(x_pt - 789) < 10 and abs(y_pt - 235) < 10:
            style_card(shape, bg_color=None, border_color=PALETTE["violet"], border_width=1.5, rounded=True)
        elif abs(x_pt - 957) < 10 and abs(y_pt - 314) < 10:
            style_card(shape, bg_color=None, border_color=PALETTE["cyan"], border_width=1.5, rounded=True)
        elif y_pt >= 420 and y_pt <= 435 and shape.has_text_frame:
            text = shape.text.strip()
            if text in {"01", "02", "03"}:
                style_card(shape, bg_color="EEF0F8", border_color=PALETTE["violet"], border_width=1, rounded=True)
                set_text(shape, text, font_name="Aptos", size=11, color=PALETTE["violet"], bold=True, alignment=PP_ALIGN.CENTER)

    # index 10: Cambio de Régimen
    slide = prs.slides[10]
    for shape in slide.shapes:
        x_pt = shape.left / 12700
        y_pt = shape.top / 12700
        w_pt = shape.width / 12700
        if abs(y_pt - 340) < 10 and w_pt > 400:
            style_card(shape, border_color=PALETTE["cyan"], border_width=1.5)

    # index 12: Intención Latente (formerly 11)
    slide = prs.slides[12]
    for shape in slide.shapes:
        x_pt = shape.left / 12700
        y_pt = shape.top / 12700
        w_pt = shape.width / 12700
        h_pt = shape.height / 12700
        if abs(w_pt - 228) < 10 and abs(h_pt - 82) < 10:
            if abs(x_pt - 560) < 10 and abs(y_pt - 340) < 10:
                style_card(shape, bg_color="EBE8FF", border_color=PALETTE["violet"], border_width=1.5, rounded=True)
            elif abs(x_pt - 820) < 10 and abs(y_pt - 340) < 10:
                style_card(shape, bg_color="F0EEFF", border_color=PALETTE["violet"], border_width=1.2, rounded=True)
            elif abs(x_pt - 560) < 10 and abs(y_pt - 456) < 10:
                style_card(shape, bg_color="F8F7FF", border_color=PALETTE["line_light"], border_width=1, rounded=True)
            elif abs(x_pt - 820) < 10 and abs(y_pt - 456) < 10:
                style_card(shape, bg_color="FAF9FF", border_color=PALETTE["line_light"], border_width=1, rounded=True)
        elif shape.has_text_frame:
            text = shape.text.strip()
            if text in {"40%", "35%"}:
                set_text(shape, text, font_name="Aptos Display", size=18, color=PALETTE["violet"], bold=True, alignment=PP_ALIGN.RIGHT)
            elif text in {"15%", "10%"}:
                set_text(shape, text, font_name="Aptos Display", size=18, color=PALETTE["body_light"], bold=True, alignment=PP_ALIGN.RIGHT)

    # index 13: Propagación de Incertidumbre (formerly 12)
    slide = prs.slides[13]
    for shape in slide.shapes:
        x_pt = shape.left / 12700
        y_pt = shape.top / 12700
        w_pt = shape.width / 12700
        if abs(y_pt - 418) < 10 and abs(w_pt - 80) < 5:
            if abs(x_pt - 235) < 10:
                style_card(shape, bg_color=PALETTE["violet"], border_color=None, rounded=False)
            elif abs(x_pt - 405) < 10:
                style_card(shape, bg_color=PALETTE["cyan"], border_color=None, rounded=False)
            elif abs(x_pt - 575) < 10:
                style_card(shape, bg_color=PALETTE["orange"], border_color=None, rounded=False)

    # index 14: Retrieval (Semantic) (formerly 13)
    slide = prs.slides[14]
    for shape in slide.shapes:
        x_pt = shape.left / 12700
        y_pt = shape.top / 12700
        w_pt = shape.width / 12700
        h_pt = shape.height / 12700
        if abs(w_pt - 18) < 2 and abs(h_pt - 18) < 2:
            if x_pt < 900:
                style_card(shape, bg_color=PALETTE["violet"], border_color=None, rounded=True)
            else:
                style_card(shape, bg_color=PALETTE["cyan"], border_color=None, rounded=True)
        elif abs(x_pt - 930) < 10 and abs(y_pt - 332) < 10:
            style_card(shape, bg_color=PALETTE["orange"], border_color=PALETTE["white"], border_width=1.5, rounded=True)

    # index 15: Reranking (formerly 14)
    slide = prs.slides[15]
    for shape in slide.shapes:
        x_pt = shape.left / 12700
        y_pt = shape.top / 12700
        w_pt = shape.width / 12700
        h_pt = shape.height / 12700
        if abs(y_pt - 286) < 10 and abs(w_pt - 430) < 10:
            style_card(shape, bg_color="F1F3F9", border_color=PALETTE["line_light"], border_width=1, rounded=True)
        elif abs(x_pt - 555) < 10 and abs(y_pt - 356) < 10:
            style_card(shape, bg_color="FDF0F0", border_color=PALETTE["red"], border_width=1, rounded=True)

    # index 16: Generación (formerly 15)
    slide = prs.slides[16]
    for shape in slide.shapes:
        x_pt = shape.left / 12700
        y_pt = shape.top / 12700
        w_pt = shape.width / 12700
        h_pt = shape.height / 12700
        if abs(x_pt - 72) < 10 and abs(w_pt - 278) < 10:
            style_card(shape, bg_color="F1F3F9", border_color=PALETTE["line_light"], border_width=1, rounded=True)
        elif abs(x_pt - 610) < 10 and abs(y_pt - 326) < 10:
            style_card(shape, bg_color="F0EEFF", border_color=PALETTE["violet"], border_width=1.5, rounded=True)
        elif abs(x_pt - 1050) < 10 and abs(y_pt - 341) < 10:
            style_card(shape, bg_color="FDF0F0", border_color=PALETTE["red"], border_width=1.5, rounded=True)

    # index 17: Límite Teórico (formerly 16)
    slide = prs.slides[17]
    for shape in slide.shapes:
        x_pt = shape.left / 12700
        y_pt = shape.top / 12700
        w_pt = shape.width / 12700
        h_pt = shape.height / 12700
        if abs(x_pt - 500) < 10 and abs(y_pt - 208) < 10:
            style_card(shape, bg_color="FDF0F0", border_color=PALETTE["red"], border_width=1.5, rounded=True)

    # index 18: Clarificación (formerly 17)
    slide = prs.slides[18]
    for shape in slide.shapes:
        x_pt = shape.left / 12700
        y_pt = shape.top / 12700
        w_pt = shape.width / 12700
        h_pt = shape.height / 12700
        if abs(x_pt - 72) < 10 and abs(y_pt - 238) < 10:
            style_card(shape, bg_color="F1F3F9", border_color=PALETTE["line_light"], border_width=1, rounded=True)
        elif abs(x_pt - 700) < 10 and abs(y_pt - 410) < 10:
            style_card(shape, bg_color="E8FCFA", border_color=PALETTE["cyan"], border_width=1.5, rounded=True)

    # index 19: Política de Interacción (formerly 18)
    slide = prs.slides[19]
    for shape in slide.shapes:
        x_pt = shape.left / 12700
        y_pt = shape.top / 12700
        w_pt = shape.width / 12700
        h_pt = shape.height / 12700
        if abs(y_pt - 302) < 10 and abs(w_pt - 475) < 10:
            if abs(x_pt - 95) < 10:
                style_card(shape, bg_color="ECFDF3", border_color=PALETTE["green"], border_width=1.5, rounded=True)
            elif abs(x_pt - 710) < 10:
                style_card(shape, bg_color="F0EEFF", border_color=PALETTE["violet"], border_width=1.5, rounded=True)

    # index 20: Evidencia (formerly 19)
    slide = prs.slides[20]
    for shape in slide.shapes:
        x_pt = shape.left / 12700
        y_pt = shape.top / 12700
        w_pt = shape.width / 12700
        h_pt = shape.height / 12700
        if abs(x_pt - 72) < 10 and abs(y_pt - 232) < 10:
            style_card(shape, bg_color="F1F3F9", border_color=PALETTE["line_light"], border_width=1, rounded=True)

    # Add a background card for statistics on the right side of Slide 21 (index 20)
    try:
        bg_card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(8.0), Inches(2.3), Inches(3.8), Inches(4.16))
        style_card(bg_card, bg_color="F8F9FC", border_color=PALETTE["line_light"], border_width=1, rounded=True)
        # Send behind textboxes
        slide.shapes._spTree.remove(bg_card._element)
        slide.shapes._spTree.insert(2, bg_card._element)
    except Exception:
        pass

    for shape in slide.shapes:
        if shape.has_text_frame:
            text = shape.text.strip()
            if text == "54% → 67%":
                set_text(shape, "54% → 67%", font_name="Aptos Display", size=24, color=PALETTE["violet"], bold=True, alignment=PP_ALIGN.LEFT)
            elif text == "46% → 55%":
                set_text(shape, "46% → 55%", font_name="Aptos Display", size=24, color=PALETTE["red"], bold=True, alignment=PP_ALIGN.LEFT)
            elif text == "0–2.5%":
                set_text(shape, "0–2.5%", font_name="Aptos Display", size=24, color=PALETTE["orange"], bold=True, alignment=PP_ALIGN.LEFT)

    # index 21: Blueprint Operativo (formerly 20)
    slide = prs.slides[21]
    for shape in slide.shapes:
        x_pt = shape.left / 12700
        y_pt = shape.top / 12700
        w_pt = shape.width / 12700
        h_pt = shape.height / 12700
        if abs(y_pt - 281) < 10 and abs(h_pt - 132) < 10 and abs(w_pt - 190) < 10:
            style_card(shape, bg_color="F1F3F9", border_color=PALETTE["line_light"], border_width=1, rounded=True)
        elif shape.has_text_frame:
            text = normalized_text(shape.text)
            if "{Caso de uso}" in text and "{Política de riesgo}" in text:
                set_text(shape, "Caso de Negocio  ·  Tolerancia al Error  ·  Mapeo de Intención", 
                         font_name="Aptos", size=14, color=PALETTE["body_dark"], bold=False, alignment=PP_ALIGN.CENTER)


def validate(prs: Presentation) -> None:
    if len(prs.slides) != 22:
        raise ValueError(f"Se esperaban 22 diapositivas y se encontraron {len(prs.slides)}")

    unresolved = []
    for slide_no, slide in enumerate(prs.slides, 1):
        for shape in iter_text_shapes(slide):
            text = normalized_text(shape.text)
            if "Pθ(" in text or "Attention(" in text or "Raclarar" in text or "A* = arg" in text:
                unresolved.append((slide_no, text))
    if unresolved:
        raise ValueError(f"Quedaron fórmulas antiguas sin convertir: {unresolved}")


def improve(input_path: Path, output_path: Path) -> None:
    prs = Presentation(str(input_path))
    if len(prs.slides) != 20:
        raise ValueError(f"La entrada debe tener 20 diapositivas; tiene {len(prs.slides)}")

    with tempfile.TemporaryDirectory(prefix="context_equations_") as temp:
        tmp_dir = Path(temp)
        update_equations(prs, tmp_dir)
        add_new_comparison_slide(prs, tmp_dir)
        add_conversational_challenges_slide(prs)

    apply_styling_improvements(prs)
    renumber_slides(prs)
    validate(prs)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    print(f"PPTX mejorada: {output_path}")
    print(f"Diapositivas: {len(prs.slides)}")
    print("Ecuaciones: revisadas, corregidas y renderizadas con notación matemática")



def parse_args():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--_reexec", action="store_true", help=argparse.SUPPRESS)
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    improve(args.input.resolve(), args.output.resolve())
